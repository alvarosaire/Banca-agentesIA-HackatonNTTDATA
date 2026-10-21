"""Extracción de texto por página: PDF (pypdf, respaldo pdfplumber), DOCX (python-docx) y TXT.

    paginas, requiere_revision, motivo = parsear(path)
    paginas = [(1, "texto página 1"), (2, "...")]
"""
from __future__ import annotations

import logging
import re
from pathlib import Path

log = logging.getLogger("parsing")

# Encabezados de artículo, también en Markdown ("## Artículo 2.", "**Art. 3**")
ART_RE = re.compile(r"^[ \t]*(?:#{1,6}[ \t]*)?(?:\*\*)?[ \t]*(Art[íi]culo|Art\.)[ \t]+\d+", re.IGNORECASE | re.MULTILINE)
DISP_RE = re.compile(r"^[ \t]*(?:#{1,6}[ \t]*)?(?:\*\*)?[ \t]*Disposici[oó]n(?:es)?[ \t]+"
                     r"(?:complementarias?[ \t]+)?finale?s?", re.IGNORECASE | re.MULTILINE)
FM_RE = re.compile(r"\A﻿?\s*---[ \t]*\r?\n(.*?)\r?\n---[ \t]*(?:\r?\n|\Z)", re.S)


# Ficha de metadatos de los PDF/DOCX: etiqueta en una línea y valor en la(s) siguiente(s)
FICHA = {"código": "codigo", "codigo": "codigo", "entidad emisora": "entidad_emisora",
         "fecha de emisión": "fecha_emision", "fecha de emision": "fecha_emision",
         "fecha de vigencia": "fecha_vigencia", "estado": "estado", "versión": "version", "version": "version",
         "reemplaza a": "reemplaza_a",
         "áreas de aplicación": "areas_aplica", "areas de aplicacion": "areas_aplica",
         "áreas de aplicacion": "areas_aplica", "areas de aplicación": "areas_aplica"}
# Variante en línea: "Áreas de aplicación: Marketing, Legal"
AREAS_INLINE_RE = re.compile(r"^[áa]reas\s+de\s+aplicaci[óo]n\s*:\s*(.+)$", re.IGNORECASE)
SIN_VALOR = {"", "ninguna", "ninguno", "no aplica", "n/a", "-", "—"}
# Encabezado/pie que se repite en cada página del PDF ("RES-SINT-2024-001 — v2.0", "Página 2")
CABECERA_RE = re.compile(r"^[ 	]*(?:[A-Z][A-Z0-9]*(?:-[A-Z0-9]+)+[ 	]+[—-][ 	]+v[\d.]+|P[áa]gina[ 	]+\d+(?:[ 	]+de[ 	]+\d+)?)[ 	]*$",
                         re.MULTILINE)


def limpiar_cabeceras(texto: str) -> str:
    return CABECERA_RE.sub("", texto or "")


def separar_ficha(texto: str) -> tuple[dict[str, str | None], str]:
    """Ficha 'Código / RES-…', 'Versión / 2.0'… antes del primer artículo -> (dict, cuerpo).

    Solo lee lo declarado en el documento (nada se infiere aquí). Sin ficha reconocible -> ({}, texto).
    """
    m_art = ART_RE.search(texto or "")
    cabeza = texto[: m_art.start()] if m_art else (texto or "")[:3000]
    lineas = [l.strip() for l in cabeza.splitlines()]
    inline = None
    for i, l in enumerate(lineas):
        m_in = AREAS_INLINE_RE.match(l)
        if m_in:
            inline = m_in.group(1).strip()
            lineas[i] = ""
    idx = [i for i, l in enumerate(lineas) if l.lower() in FICHA]
    if "codigo" not in {FICHA[lineas[i].lower()] for i in idx}:
        return {}, texto
    fm: dict[str, str | None] = {}
    for n, i in enumerate(idx):
        fin = idx[n + 1] if n + 1 < len(idx) else len(lineas)
        valor = []
        for l in lineas[i + 1: fin]:
            if not l or l.lower().startswith("aviso"):
                break
            valor.append(l)
        v = " ".join(valor).strip()
        fm[FICHA[lineas[i].lower()]] = None if v.lower() in SIN_VALOR else v
    if inline and not fm.get("areas_aplica"):
        fm["areas_aplica"] = inline
    nombre = " ".join(l for l in lineas[: idx[0]] if l).strip()
    if nombre:
        fm["nombre_norma"] = nombre
    return fm, texto[m_art.start():] if m_art else texto


def separar_frontmatter(texto: str) -> tuple[dict[str, str | None], str]:
    """Frontmatter YAML simple (clave: valor) al inicio del texto -> (dict, cuerpo). Valor vacío = None.

    Sin YAML, se intenta la ficha de metadatos de los PDF/DOCX (separar_ficha).
    """
    m = FM_RE.match(texto or "")
    if not m:
        return separar_ficha(texto)
    fm: dict[str, str | None] = {}
    for linea in m.group(1).splitlines():
        if ":" not in linea or linea.lstrip().startswith("#"):
            continue
        k, v = linea.split(":", 1)
        v = v.strip().strip('"').strip("'").strip()
        fm[k.strip().lower()] = v or None
    return fm, texto[m.end():]


def _pdf(path: Path) -> list[tuple[int, str]]:
    try:
        from pypdf import PdfReader

        pages = [(i + 1, limpiar_cabeceras(p.extract_text() or "")) for i, p in enumerate(PdfReader(str(path)).pages)]
        if sum(len(t.strip()) for _, t in pages) > 50:
            return pages
        log.warning("pypdf extrajo poco texto de %s; se intenta pdfplumber", path.name)
    except Exception as e:
        log.warning("pypdf falló con %s: %s; se intenta pdfplumber", path.name, e)
    import pdfplumber

    with pdfplumber.open(str(path)) as pdf:
        return [(i + 1, limpiar_cabeceras(p.extract_text() or "")) for i, p in enumerate(pdf.pages)]


def _docx(path: Path) -> list[tuple[int, str]]:
    import docx

    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return [(1, "\n".join(lines))]


def _txt(path: Path) -> list[tuple[int, str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    # \f (salto de página) separa páginas si existe
    return [(i + 1, t) for i, t in enumerate(text.split("\f"))]


def es_basura(texto: str) -> tuple[bool, str]:
    """Heurística: >30% de caracteres no alfabéticos (sin contar espacios) o sin artículos."""
    sin_esp = re.sub(r"\s", "", texto)
    if len(sin_esp) < 50:
        return True, "Texto vacío o casi vacío (¿PDF escaneado sin OCR?)"
    sin_esp = re.sub(r"[|#*>\-]", "", sin_esp)  # marcas de Markdown/tablas no cuentan como basura
    if not sin_esp:
        return True, "Texto vacío o casi vacío"
    no_alfa = sum(1 for c in sin_esp if not c.isalpha())
    if no_alfa / len(sin_esp) > 0.30:
        return True, f"{no_alfa / len(sin_esp):.0%} de caracteres no alfabéticos"
    if not ART_RE.search(texto) and not DISP_RE.search(texto):
        return True, "No se detectaron artículos numerados"
    return False, ""


def parsear(path: str | Path) -> tuple[list[tuple[int, str]], bool, str]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        pages = _pdf(path)
    elif ext == ".docx":
        pages = _docx(path)
    elif ext in (".txt", ".md", ""):
        pages = _txt(path)
    else:
        raise ValueError(f"Formato no soportado: {ext} (usa PDF, DOCX, TXT o MD)")
    _, cuerpo = separar_frontmatter("\n".join(t for _, t in pages))
    basura, motivo = es_basura(cuerpo)
    return pages, basura, motivo
